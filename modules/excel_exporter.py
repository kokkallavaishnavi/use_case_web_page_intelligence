"""
Excel Exporter Module
====================

This module handles exporting extracted Golden Questions and Expected Responses
to Excel format for bot testing purposes.

Author: AI Assistant
Date: 2025
"""

import os
import logging
from typing import List, Dict, Any
from datetime import datetime
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

from config.settings import AppConfig
from utils.helpers import sanitize_filename, truncate_text

logger = logging.getLogger(__name__)

class ExcelExporter:
    """Export extracted questions to Excel format for bot testing"""
    
    def __init__(self):
        """Initialize the Excel exporter"""
        self.config = AppConfig()
        self.export_stats = {
            "total_exports": 0,
            "successful_exports": 0,
            "failed_exports": 0,
            "total_questions_exported": 0
        }
    
    def _create_filename(self, session_id: str = None, custom_name: str = None) -> str:
        """Create filename for Excel export"""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        
        if custom_name:
            base_name = sanitize_filename(custom_name)
        elif session_id:
            base_name = f"bot_testing_questions_{session_id}"
        else:
            base_name = "bot_testing_questions"
        
        filename = f"{base_name}_{timestamp}.xlsx"
        return os.path.join(self.config.EXPORT_DIR, filename)
    
    def _prepare_questions_data(self, questions: List[Dict[str, Any]]) -> pd.DataFrame:
        """Prepare questions data for Excel export"""
        data_rows = []
        
        for i, question in enumerate(questions, 1):
            question_text = truncate_text(
                question.get("question", ""), 
                self.config.MAX_CELL_LENGTH
            )
            response_text = truncate_text(
                question.get("expected_response", ""), 
                self.config.MAX_CELL_LENGTH
            )
            
            row = {
                "Test_Case_ID": f"TC_{i:03d}",
                "Question": question_text,
                "Expected_Response": response_text,
                "Keywords": ", ".join(question.get("keywords", [])),
                "Source_URL": question.get("source_url", ""),
            }
            data_rows.append(row)
        
        return pd.DataFrame(data_rows)
    
    def _style_worksheet(self, worksheet, df: pd.DataFrame):
        """Apply styling to the worksheet"""
        try:
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            
            border_style = Side(style="thin", color="000000")
            border = Border(left=border_style, right=border_style, top=border_style, bottom=border_style)
            
            for cell in worksheet[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
                cell.border = border
            
            column_widths = {
                'A': 12,
                'B': 50,
                'C': 60,
                'D': 30,
                'E': 40,
            }
            for col_letter, width in column_widths.items():
                worksheet.column_dimensions[col_letter].width = width
                
            for row in range(2, len(df) + 2):
                for col in range(1, len(df.columns) + 1):
                    cell = worksheet.cell(row=row, column=col)
                    cell.border = border
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
            
            worksheet.freeze_panes = "A2"
            logger.info("Worksheet styling applied successfully")
        except Exception as e:
            logger.warning(f"Error applying worksheet styling: {str(e)}")
    
    def export_questions(self, questions: List[Dict[str, Any]], 
                        session_id: str = None, 
                        custom_filename: str = None) -> str:
        """Export questions to Excel file"""
        self.export_stats["total_exports"] += 1
        
        try:
            if not questions:
                raise ValueError("No questions provided for export")
            
            logger.info(f"Starting Excel export of {len(questions)} questions")
            
            filepath = self._create_filename(session_id, custom_filename)
            df = self._prepare_questions_data(questions)
            
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = self.config.EXCEL_SHEET_NAME
            
            for r in dataframe_to_rows(df, index=False, header=True):
                worksheet.append(r)
            
            self._style_worksheet(worksheet, df)
            
            workbook.save(filepath)
            
            self.export_stats["successful_exports"] += 1
            self.export_stats["total_questions_exported"] += len(questions)
            
            logger.info(f"Excel export completed successfully: {filepath}")
            
            return filepath
        except Exception as e:
            self.export_stats["failed_exports"] += 1
            logger.error(f"Error exporting to Excel: {str(e)}")
            raise
    
    def export_questions_with_metadata(self, questions: List[Dict[str, Any]], 
                                     metadata: Dict[str, Any],
                                     session_id: str = None) -> str:
        try:
            enhanced_questions = []
            for question in questions:
                enhanced_question = question.copy()
                enhanced_question.update({
                    "session_id": session_id or "unknown",
                    "export_timestamp": datetime.now().isoformat(),
                    **metadata
                })
                enhanced_questions.append(enhanced_question)
            
            return self.export_questions(enhanced_questions, session_id)
        except Exception as e:
            logger.error(f"Error exporting with metadata: {str(e)}")
            raise
    
    def create_test_template(self, filepath: str = None) -> str:
        try:
            if not filepath:
                filepath = self._create_filename(custom_name="test_template")
            
            template_data = pd.DataFrame(columns=[
                "Test_Case_ID", "Question", "Expected_Response", "Keywords", "Source_URL"
            ])
            
            sample_row = {
                "Test_Case_ID": "TC_001",
                "Question": "What services do you provide?",
                "Expected_Response": "We provide comprehensive bot testing services including...",
                "Keywords": "services, bot testing, comprehensive",
                "Source_URL": "https://example.com"
            }
            
            template_data = pd.concat([template_data, pd.DataFrame([sample_row])], ignore_index=True)
            
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "Test_Cases_Template"
            
            for r in dataframe_to_rows(template_data, index=False, header=True):
                worksheet.append(r)
            
            self._style_worksheet(worksheet, template_data)
            workbook.save(filepath)
            
            logger.info(f"Test template created: {filepath}")
            return filepath
        except Exception as e:
            logger.error(f"Error creating test template: {str(e)}")
            raise
    
    def get_export_stats(self) -> Dict[str, Any]:
        stats = self.export_stats.copy()
        if stats["total_exports"] > 0:
            stats["success_rate"] = stats["successful_exports"] / stats["total_exports"]
            stats["avg_questions_per_export"] = stats["total_questions_exported"] / stats["successful_exports"] if stats["successful_exports"] > 0 else 0
        else:
            stats["success_rate"] = 0.0
            stats["avg_questions_per_export"] = 0.0
        
        return stats
